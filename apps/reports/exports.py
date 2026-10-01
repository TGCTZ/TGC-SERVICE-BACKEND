"""Generate complete filtered downloads from the same projections as the API."""

from decimal import Decimal
from io import BytesIO
from math import ceil
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from django.utils import timezone

from apps.core.exceptions import ServiceError

from .definitions import EXPORT_ROW_LIMIT, REPORT_TIMEZONE
from .selectors import report_columns, report_rows


def _metadata(report):
    return [
        ("Report", report["title"]),
        ("From", report["range"]["from"]),
        ("To", report["range"]["to"]),
        ("Timezone", str(REPORT_TIMEZONE)),
        (
            "Generated at",
            timezone.localtime(report["generated_at"], REPORT_TIMEZONE).isoformat(),
        ),
        *[(name, str(value)) for name, value in report["applied_filters"].items()],
    ]


def _summary_lines(summary):
    lines = [("Matching records", summary["count"])]
    lines += [
        (f"Total ({row['currency']})", Decimal(row["amount"]))
        for row in summary["amounts"]
    ]
    lines += [
        (f"Current status: {row['status']}", row["count"]) for row in summary["statuses"]
    ]
    lines.append(
        (
            "Undated records excluded (all dates, matching other filters)",
            summary["missing_dates"],
        )
    )
    return lines


def _xlsx(report, sections):
    workbook = Workbook(write_only=True)
    for key, section in sections.items():
        summary = section["summary"]
        sheet = workbook.create_sheet(summary["title"][:31])
        metadata = [
            *_metadata(report),
            ("Description", summary["description"]),
            *_summary_lines(summary),
        ]
        columns = report_columns(key)
        sheet.freeze_panes = f"A{len(metadata) + 3}"
        widths = [
            40
            if column["key"] == "customer"
            else (24 if column["kind"] == "text" else 18)
            for column in columns
        ]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row_number, (name, value) in enumerate(metadata, 1):
            sheet.row_dimensions[row_number].height = 15 * max(
                ceil(len(str(name)) / widths[0]), ceil(len(str(value)) / widths[1])
            )
            # Force strings to remain text, including customer input beginning '='.
            cells = []
            for item in (name, value):
                cell = WriteOnlyCell(sheet, value=item)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(item, str):
                    cell.data_type = "s"
                elif isinstance(item, Decimal):
                    cell.number_format = "#,##0.00"
                cells.append(cell)
            sheet.append(cells)
        sheet.append([])
        headers = []
        for column in columns:
            cell = WriteOnlyCell(sheet, value=column["label"])
            cell.font = Font(bold=True)
            headers.append(cell)
        sheet.append(headers)
        for row_number, row in enumerate(
            report_rows(key, section["queryset"].iterator(chunk_size=500)),
            len(metadata) + 3,
        ):
            sheet.row_dimensions[row_number].height = 15 * max(
                ceil(len(str(row[column["key"]])) / width)
                for column, width in zip(columns, widths, strict=True)
            )
            cells = []
            for column in columns:
                value = row[column["key"]]
                cell = WriteOnlyCell(sheet, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if column["kind"] == "money":
                    cell.value = float(value)
                    cell.number_format = "#,##0.00"
                elif isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(doc.pagesize[0] - 28, 16, f"Page {doc.page}")
    canvas.restoreState()


def _pdf(report, sections):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=28,
    )
    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 10
    money_style = ParagraphStyle(
        "ReportMoney", parent=styles["BodyText"], alignment=TA_RIGHT
    )
    story = [Paragraph(escape(report["title"]), styles["Title"])]
    for name, value in _metadata(report):
        story.append(Paragraph(escape(f"{name}: {value}"), styles["BodyText"]))
    for key, section in sections.items():
        summary = section["summary"]
        story += [
            Spacer(1, 16),
            Paragraph(escape(summary["title"]), styles["Heading2"]),
            Paragraph(escape(summary["description"]), styles["BodyText"]),
        ]
        for name, value in _summary_lines(summary):
            if isinstance(value, Decimal):
                value = f"{value:,.2f}"
            story.append(Paragraph(escape(f"{name}: {value}"), styles["BodyText"]))
        story.append(Spacer(1, 8))
        columns = report_columns(key)
        data = [
            [Paragraph(escape(column["label"]), styles["BodyText"]) for column in columns]
        ]
        for row in report_rows(key, section["queryset"].iterator(chunk_size=500)):
            data.append(
                [
                    Paragraph(escape(f"{Decimal(row[column['key']]):,.2f}"), money_style)
                    if column["kind"] == "money"
                    else Paragraph(escape(str(row[column["key"]])), styles["BodyText"])
                    for column in columns
                ]
            )
        if len(data) == 1:
            story.append(
                Paragraph("Nothing recorded in this period.", styles["BodyText"])
            )
        else:
            table = LongTable(
                data,
                colWidths=[document.width / len(columns)] * len(columns),
                repeatRows=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7eef0")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#ccd5d8")),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.append(table)
    document.build(story, onFirstPage=_page_number, onLaterPages=_page_number)
    return output.getvalue()


def export_report(report: dict, sections: dict, *, file_type: str) -> bytes:
    """Reject oversized results before fetching rows; downloads never truncate."""
    if (
        sum(section["summary"]["count"] for section in sections.values())
        > EXPORT_ROW_LIMIT
    ):
        raise ServiceError(
            f"Exports allow at most {EXPORT_ROW_LIMIT:,} rows. Narrow the report filters."
        )
    return _xlsx(report, sections) if file_type == "xlsx" else _pdf(report, sections)
