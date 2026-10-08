"""Report dates, monetary invariants, source permissions and complete exports."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from django.contrib.auth.models import Permission
from django.db import connection
from django.test.utils import CaptureQueriesContext

from apps.billing.models import Bill, Payment
from apps.certificates.models import Certificate
from apps.core.current_user import reset_current_user, set_current_user
from apps.identification.models import IdentificationReport
from apps.orders.tests.factories import OrderFactory, StoneFactory

pytestmark = pytest.mark.django_db
LAB = ZoneInfo("Africa/Dar_es_Salaam")
DATES = {"from_date": "2026-09-01", "to_date": "2026-09-30"}


def test_schema_documents_report_filters_and_binary_exports():
    """Keep the report's public contract discoverable by generated clients."""
    from drf_spectacular.generators import SchemaGenerator

    schema = SchemaGenerator().get_schema(request=None, public=True)
    result = schema["paths"]["/api/v1/reports/financial/"]["get"]
    assert {"from_date", "to_date", "section", "page_size"}.issubset(
        {parameter["name"] for parameter in result["parameters"]}
    )
    content = schema["paths"]["/api/v1/reports/financial/export/"]["get"]["responses"][
        "200"
    ]["content"]
    assert "application/pdf" in content
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in content


def _at(day, hour=12):
    return datetime(2026, 9, day, hour, tzinfo=LAB)


def _grant(user, *names):
    for name in names:
        app, codename = name.split(".")
        user.user_permissions.add(
            Permission.objects.get(content_type__app_label=app, codename=codename)
        )


def _bill(**fields):
    order = fields.pop("order", None) or OrderFactory(received_date="2026-09-01")
    return Bill.objects.create(
        order=order,
        bill_number=f"BILL-{order.pk}",
        issued_at=_at(10),
        total_amount=Decimal("100.00"),
        **fields,
    )


def _payment(bill, **fields):
    values = {
        "bill": bill,
        "trx_id": f"TRX-{Payment.objects.count()}",
        "trx_dt_tm": _at(15),
        "paid_amount": Decimal("25.00"),
        "currency": bill.currency,
        "is_processed": True,
    }
    return Payment.objects.create(**(values | fields))


def _summary(response, key):
    return next(section for section in response.data["sections"] if section["key"] == key)


def test_unauthenticated_and_users_without_report_permissions_are_refused(
    api_client, user
):
    """A hidden sidebar is never the security boundary."""
    assert api_client.get("/api/v1/reports/financial/").status_code == 401
    api_client.force_authenticate(user)
    assert api_client.get("/api/v1/reports/financial/").status_code == 403
    assert api_client.get("/api/v1/reports/operational/export/").status_code == 403
    assert api_client.get("/api/v1/reports/").data == []


def test_report_gates_control_catalog_results_and_exports(api_client, user):
    """Source access alone must not bypass either report gate."""
    _grant(user, "billing.view_bill")
    api_client.force_authenticate(user)
    assert api_client.get("/api/v1/reports/").data == []
    assert api_client.get("/api/v1/reports/financial/", DATES).status_code == 403

    _grant(user, "core.module_reports")
    user = type(user).objects.get(pk=user.pk)
    api_client.force_authenticate(user)
    assert api_client.get("/api/v1/reports/financial/export/", DATES).status_code == 403

    _grant(user, "core.report_financial")
    user = type(user).objects.get(pk=user.pk)
    api_client.force_authenticate(user)
    assert [page["key"] for page in api_client.get("/api/v1/reports/").data] == [
        "financial"
    ]
    assert api_client.get("/api/v1/reports/financial/", DATES).status_code == 200


def test_bill_only_user_cannot_read_payment_totals_or_outstanding(api_client, user):
    """Combined pages must omit sections requiring another model's permission."""
    _grant(user, "core.module_reports", "core.report_financial", "billing.view_bill")
    api_client.force_authenticate(user)
    _payment(_bill())
    response = api_client.get("/api/v1/reports/financial/", DATES)
    assert response.status_code == 200
    assert [section["key"] for section in response.data["sections"]] == ["billing"]
    assert (
        api_client.get(
            "/api/v1/reports/financial/", DATES | {"section": "outstanding"}
        ).status_code
        == 403
    )
    export = api_client.get("/api/v1/reports/financial/export/", DATES)
    workbook = load_workbook(BytesIO(export.content))
    assert workbook.sheetnames == ["Billing summary"]


def test_catalog_and_operational_sections_follow_individual_permissions(api_client, user):
    """Order readers cannot discover certificate or findings data."""
    _grant(user, "core.module_reports", "core.report_operational", "orders.view_order")
    api_client.force_authenticate(user)
    catalog = api_client.get("/api/v1/reports/")
    assert [page["key"] for page in catalog.data] == ["operational"]
    response = api_client.get("/api/v1/reports/operational/", DATES)
    assert [section["key"] for section in response.data["sections"]] == ["orders"]
    assert (
        api_client.get(
            "/api/v1/reports/operational/", DATES | {"section": "certificates"}
        ).status_code
        == 403
    )


def test_totals_span_all_pages_and_currencies_without_multiplying_payments(
    api_client, admin_user
):
    """More detail rows or repeated relationships must not inflate money."""
    bill = _bill()
    StoneFactory.create_batch(3, order=bill.order)
    _payment(bill)
    _payment(bill)
    other = _bill(currency="USD")
    _payment(other, paid_amount=Decimal("10.00"))
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/financial/", DATES | {"section": "outstanding", "page_size": 1}
    )
    assert response.status_code == 200
    assert response.data["count"] == 2
    assert len(response.data["results"]) == 1
    assert _summary(response, "billing")["count"] == 2
    amounts = {
        item["currency"]: Decimal(item["amount"])
        for item in _summary(response, "outstanding")["amounts"]
    }
    assert amounts == {"TZS": Decimal("50"), "USD": Decimal("90")}


def test_outstanding_filters_issue_date_but_uses_payments_received_later(
    api_client, admin_user
):
    """Outstanding reports show today's balance for the selected bill cohort."""
    bill = _bill()
    _payment(bill, trx_dt_tm=datetime(2026, 10, 1, tzinfo=LAB), paid_amount=Decimal("60"))
    _payment(bill, is_processed=False, paid_amount=Decimal("40"))
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/financial/", DATES | {"section": "outstanding"}
    )
    assert Decimal(response.data["results"][0]["balance"]) == Decimal("40")
    assert _summary(response, "collections")["count"] == 0
    assert _summary(response, "exceptions")["count"] == 1


def test_cancelled_overpaid_and_deleted_bills_are_not_outstanding(api_client, admin_user):
    """Cancellation, overpayment and soft deletion cannot produce a chase balance."""
    _bill(status="cancelled")
    overpaid = _bill()
    _payment(overpaid, paid_amount=Decimal("120"))
    deleted = _bill()
    deleted.delete()
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/financial/", DATES | {"section": "outstanding"}
    )
    assert response.data["count"] == 0


def test_lab_day_boundaries_and_undated_notifications(api_client, admin_user):
    """UTC storage must not move transactions across lab calendar days."""
    bill = _bill()
    _payment(bill, trx_dt_tm=datetime(2026, 8, 31, 21, tzinfo=ZoneInfo("UTC")))
    _payment(bill, trx_dt_tm=datetime(2026, 9, 30, 21, tzinfo=ZoneInfo("UTC")))
    _payment(bill, trx_dt_tm=None)
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/financial/", DATES | {"section": "collections"}
    )
    assert response.data["count"] == 1
    assert response.data["results"][0]["event_date"] == "2026-09-01"
    assert _summary(response, "collections")["missing_dates"] == 1


def test_operational_events_use_independent_dates_and_distinct_orders(
    api_client, admin_user
):
    """Submitted stone counts and registration events are different measurements."""
    order = OrderFactory(received_date="2026-09-05", stone_count=7)
    stone = StoneFactory(order=order)
    second = StoneFactory(order=order, stone_type=stone.stone_type)
    type(stone).objects.filter(pk__in=[stone.pk, second.pk]).update(created_at=_at(8))
    findings = IdentificationReport.objects.create(
        stone=stone, report_number="RPT-1", is_finalized=True, identified_at=_at(20)
    )
    Certificate.objects.create(
        stone=stone,
        report=findings,
        certificate_number="CERT-1",
        weight_snapshot=Decimal("1"),
        stone_type_snapshot=stone.stone_type.name,
        issued_at=datetime(2026, 10, 1, tzinfo=LAB),
    )
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/operational/", DATES | {"stone_type": stone.stone_type_id}
    )
    assert response.status_code == 200
    counts = {section["key"]: section["count"] for section in response.data["sections"]}
    assert counts == {"orders": 1, "stones": 2, "findings": 1, "certificates": 0}
    assert response.data["results"][0]["stones"] == 7


def test_revoked_certificate_still_counts_as_issuance(api_client, admin_user):
    """Current validity must not erase an issuance event from the period."""
    stone = StoneFactory()
    findings = IdentificationReport.objects.create(stone=stone, report_number="RPT-2")
    Certificate.objects.create(
        stone=stone,
        report=findings,
        certificate_number="CERT-2",
        weight_snapshot=Decimal("1"),
        stone_type_snapshot="Quartz",
        issued_at=_at(25),
        status="revoked",
    )
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/operational/", DATES | {"section": "certificates"}
    )
    assert response.data["count"] == 1
    assert response.data["results"][0]["status"] == "Revoked"


def test_deleted_matching_stone_cannot_qualify_an_order(api_client, admin_user):
    """The type and live-record predicates must apply to the same joined stone."""
    order = OrderFactory(received_date="2026-09-05")
    removed = StoneFactory(order=order)
    StoneFactory(order=order)
    removed.delete()
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/operational/", DATES | {"stone_type": removed.stone_type_id}
    )
    assert response.status_code == 200
    assert _summary(response, "orders")["count"] == 0


@pytest.mark.parametrize(
    "filters",
    [
        {"from_date": "2026-10-01"},
        {"to_date": "invalid"},
        {"page": 0},
        {"page_size": 101},
        {"file_type": "csv"},
    ],
)
def test_invalid_filters_return_validation_errors(api_client, admin_user, filters):
    """Malformed requests do not silently produce a different report."""
    api_client.force_authenticate(admin_user)
    assert (
        api_client.get("/api/v1/reports/financial/", DATES | filters).status_code == 400
    )


def test_financial_rows_use_control_numbers_without_exposing_customer_data(
    api_client, user
):
    """Financial reports identify bills by gateway number, not customer identity."""
    _grant(
        user,
        "core.module_reports",
        "core.report_financial",
        "billing.view_bill",
        "billing.view_payment",
    )
    visible = _bill(control_number="991234567890")
    _payment(visible)
    api_client.force_authenticate(user)
    response = api_client.get("/api/v1/reports/financial/", DATES)
    assert response.status_code == 200
    assert response.data["columns"][2] == {
        "key": "control_number",
        "label": "Control number",
        "kind": "text",
    }
    assert response.data["results"][0]["control_number"] == visible.control_number
    assert "customer" not in response.data["results"][0]
    assert response.data["filters"]["customers"] == []

    collections = api_client.get(
        "/api/v1/reports/financial/", DATES | {"section": "collections"}
    )
    assert "provider" not in {column["key"] for column in collections.data["columns"]}
    assert "provider" not in collections.data["results"][0]


def test_operational_reports_keep_customer_filter_and_column(api_client, user):
    """Customer identification remains useful for operational work tracking."""
    _grant(user, "core.module_reports", "core.report_operational", "orders.view_order")
    order = OrderFactory(received_date="2026-09-05")
    api_client.force_authenticate(user)
    response = api_client.get(
        "/api/v1/reports/operational/", DATES | {"customer": order.customer_id}
    )
    assert response.status_code == 200
    assert response.data["columns"][2] == {
        "key": "customer",
        "label": "Customer",
        "kind": "text",
    }
    assert response.data["results"][0]["customer"] == order.customer.full_name


def test_xlsx_exports_every_page_and_keeps_control_numbers_as_text(
    api_client, admin_user
):
    """Downloads are complete and gateway values remain text in Excel."""
    values = [f"=1+1+{number}" for number in range(3)]
    for value in values:
        _bill(control_number=value)
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/financial/export/", DATES | {"page_size": 1}
    )
    assert response.status_code == 200
    assert "attachment;" in response["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Billing summary"]
    assert sheet["A1"].value == "Date"
    assert sheet.max_row == 4
    assert "Matching records" not in {
        cell.value for row in sheet.iter_rows() for cell in row
    }
    matching = [cell for row in sheet for cell in row if cell.value in values]
    assert {cell.value for cell in matching} == set(values)
    assert all(cell.data_type == "s" for cell in matching)


def test_pdf_has_a_valid_document_header(api_client, admin_user):
    """A binary download is a PDF, not a JSON body renamed to .pdf."""
    _bill()
    api_client.force_authenticate(admin_user)
    response = api_client.get(
        "/api/v1/reports/financial/export/", DATES | {"file_type": "pdf"}
    )
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")


def test_export_limit_rejects_instead_of_truncating(api_client, admin_user, monkeypatch):
    """Limits count all authorized sections before writing the attachment."""
    monkeypatch.setattr("apps.reports.exports.EXPORT_ROW_LIMIT", 1)
    _bill()
    api_client.force_authenticate(admin_user)
    response = api_client.get("/api/v1/reports/financial/export/", DATES)
    assert response.status_code == 400
    assert "Narrow" in str(response.data)


def test_query_count_does_not_grow_with_rows(api_client, user):
    """Joined customer names and options must not cause per-row database reads."""
    _grant(user, "core.module_reports", "core.report_financial", "billing.view_bill")
    api_client.force_authenticate(user)
    token = set_current_user(user)
    try:
        _bill()
    finally:
        reset_current_user(token)
    api_client.get("/api/v1/reports/financial/", DATES)
    with CaptureQueriesContext(connection) as small:
        response = api_client.get("/api/v1/reports/financial/", DATES)
    token = set_current_user(user)
    try:
        for _ in range(8):
            _bill()
    finally:
        reset_current_user(token)
    with CaptureQueriesContext(connection) as large:
        more = api_client.get("/api/v1/reports/financial/", DATES)
    assert response.status_code == more.status_code == 200
    assert len(large) == len(small)
